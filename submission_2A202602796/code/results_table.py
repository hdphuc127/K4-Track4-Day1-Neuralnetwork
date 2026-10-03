"""results_table.py — lưu kết quả từng lần chạy ra JSON, rồi điền experiments.xlsx từ mẫu
templates/experiment_table_template.xlsx bằng code (không gõ tay).

Các cột công thức (step0_gap_vs_lnC, gap_val_minus_train, delta_val_f1_vs_base, beyond_noise) KHÔNG bị ghi đè.
"""
from __future__ import annotations

import json
import math
from pathlib import Path

import openpyxl

OPT_NAMES = {"sgd": "SGD", "sgd_momentum": "SGD+momentum", "adam": "Adam", "adamw": "AdamW"}
FORMULA_COLS = {"step0_gap_vs_lnC", "gap_val_minus_train", "delta_val_f1_vs_base", "beyond_noise"}
MAX_ROWS = 60   # mẫu có sẵn 60 dòng công thức (Experiments!A2:A61)


def save_result(result: dict, results_dir: str = "../results") -> str:
    """Ghi cfg, history, summary (KHÔNG ghi best_state) ra <results_dir>/<exp_id>.json."""
    Path(results_dir).mkdir(parents=True, exist_ok=True)
    path = Path(results_dir) / f"{result['cfg']['exp_id']}.json"
    out = _clean({k: result[k] for k in ("cfg", "history", "summary")})
    path.write_text(json.dumps(out, indent=1, allow_nan=False))
    return str(path)


def _clean(x):
    """tuple -> list; NaN/inf -> None (JSON hợp lệ; run bị diverged có NaN trong lịch sử)."""
    if isinstance(x, dict):
        return {k: _clean(v) for k, v in x.items()}
    if isinstance(x, (list, tuple)):
        return [_clean(v) for v in x]
    if isinstance(x, float) and not math.isfinite(x):
        return None
    return x


def load_results(results_dir: str = "../results") -> list[dict]:
    """Đọc mọi *.json trong results_dir, sắp theo exp_id. cfg['hidden'] đưa về tuple."""
    out = []
    for p in Path(results_dir).glob("*.json"):
        r = json.loads(p.read_text())
        if "cfg" not in r:      # bỏ qua file phụ (ví dụ eval_scores.json)
            continue
        r["cfg"]["hidden"] = tuple(r["cfg"]["hidden"])
        out.append(r)
    return sorted(out, key=lambda r: r["cfg"]["exp_id"])


def to_row(result: dict, eval_scores: dict | None = None, notes: str = "") -> dict:
    """Một kết quả -> một dòng bảng (khoá = tên cột). Chỉ truyền eval_scores cho baseline và cấu hình cuối cùng."""
    cfg, s = result["cfg"], result["summary"]
    row = dict(
        exp_id=cfg["exp_id"], group=cfg["group"], description=cfg["description"],
        loss=cfg["loss"].upper(), optimizer=OPT_NAMES[cfg["optimizer"]], lr=cfg["lr"],
        weight_decay=cfg["weight_decay"], batch=cfg["batch"], epochs=cfg["epochs"],
        hidden="-".join(map(str, cfg["hidden"])), dropout=cfg["dropout"],
        clip_norm=cfg["clip_norm"] if cfg["clip_norm"] else "none", precision=cfg["precision"], init=cfg["init"],
        seed=cfg["seed"], diverged="Y" if s["diverged"] else "N", figure_file=f"figures/{cfg['exp_id']}.png", notes=notes,
    )
    for k in ("step0_loss", "best_val_loss", "best_epoch", "final_train_loss", "final_val_loss", "val_acc",
              "val_macro_f1", "time_per_epoch_s", "peak_mem_MB"):
        row[k] = s[k]
    if eval_scores:
        row["eval_acc"], row["eval_macro_f1"] = eval_scores["accuracy"], eval_scores["macro_f1"]
    return row


def write_xlsx(rows: list[dict], template_path: str, out_path: str,
               seed_ids: list[str] | None = None, summary_notes: dict | None = None) -> None:
    """Điền `rows` vào sheet Experiments của mẫu (từ dòng 2), seed_ids vào Seeds!A2:A6, nhận xét vào Summary!H.

    Mở bằng Excel/LibreOffice để các công thức tính lại (openpyxl không tính công thức).
    """
    if len(rows) > MAX_ROWS:
        raise ValueError(f"mẫu chỉ có sẵn {MAX_ROWS} dòng công thức, nhận {len(rows)} dòng")
    wb = openpyxl.load_workbook(template_path)    # KHÔNG dùng data_only=True (mất công thức)
    ws = wb["Experiments"]
    col = {c.value: c.column for c in ws[1] if c.value}
    for r in range(2, 2 + MAX_ROWS):               # xoá dữ liệu mẫu cũ (dòng baseline điền sẵn), giữ công thức
        for name, c in col.items():
            if name not in FORMULA_COLS:
                ws.cell(r, c).value = None
    for i, row in enumerate(rows):
        for name, v in row.items():
            if name in FORMULA_COLS:
                continue
            ws.cell(2 + i, col[name]).value = v

    if seed_ids is not None:
        assert len(seed_ids) <= 5, "sheet Seeds có 5 dòng (A2:A6)"
        for i in range(5):
            wb["Seeds"].cell(2 + i, 1).value = seed_ids[i] if i < len(seed_ids) else None
    if summary_notes:
        sm = wb["Summary"]
        for r in range(2, 12):
            g = sm.cell(r, 1).value
            if g in summary_notes:
                sm.cell(r, 8).value = summary_notes[g]
    wb.save(out_path)
